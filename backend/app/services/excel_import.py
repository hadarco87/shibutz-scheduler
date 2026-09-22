"""Flexible Excel/CSV-ish spreadsheet import for people (כוח אדם)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from sqlalchemy.orm import Session, joinedload

from app.models import (
    LabelSelectionMode,
    Person,
    PersonLabelAssignment,
    PersonQualification,
    Qualification,
    Role,
)
from app.services.person_labels import get_or_create_option, load_company_labels


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).replace("\xa0", " ").strip().lower()
    s = s.replace('"', "").replace("'", "").replace("״", "").replace("׳", "")
    s = s.replace("־", "-").replace("–", "-")
    s = re.sub(r"\s+", " ", s)
    return s


FIELD_ALIASES: Dict[str, Tuple[str, ...]] = {
    "personal_number": (
        "מספר אישי",
        "מס אישי",
        "מספראישי",
        "תעודת זהות",
        "תעודתזהות",
        'ת"ז',
        "תז",
        "ת.ז",
        "ת.ז.",
        "id",
        "personal number",
        "personal_id",
        "soldier id",
        "service number",
    ),
    "first_name": ("שם פרטי", "פרטי", "first name", "firstname", "given name"),
    "last_name": ("שם משפחה", "משפחה", "מש", "last name", "lastname", "family name", "surname"),
    "full_name": ("שם מלא", "שם החייל", "full name", "name", "שם"),
    "phone": ("טלפון", "טל", "נייד", "phone", "mobile", "cell"),
    "pakal": ('פק"ל', "פקל", "pakal", "qualification", "qualifications", "הכשרה"),
    "role": ("תפקיד", "role", "position"),
    "platoon": ("פלוגה", "יחידה", "company", "platoon"),
    "department": ("מחלקה", "מח", "department", "platoon section"),
    "squad": ("כיתה", "squad", "team"),
    "rank": ("דרגה", "rank"),
}

SKIP_SHEET_PREFIXES = ("_", "תרשים", "מקרא")
PREFERRED_SHEET_HINTS = (
    "לוח יציאות - מתוכנן",
    "לוח יציאות",
    "כוח אדם",
    "חיילים",
    "אנשים",
)


COMMANDER_HINTS = ("מפקד", "מ״כ", 'מ"כ', "סמבצ", 'סמב"צ', "סמ״צ")
NCO_HINTS = ("משק", "מש״ק", 'מש"ק', "נגד")


@dataclass
class ParsedPersonRow:
    full_name: str
    personal_number: Optional[str] = None
    phone: Optional[str] = None
    role_name: Optional[str] = None
    qualification_names: List[str] = field(default_factory=list)
    # display label name -> option value strings
    label_values: Dict[str, List[str]] = field(default_factory=dict)
    notes: Optional[str] = None
    warnings: List[str] = field(default_factory=list)
    raw: Dict[str, Any] = field(default_factory=dict)


@dataclass
class ParseResult:
    sheet_name: str
    sheet_options: List[str]
    column_mapping: Dict[str, str]
    rows: List[ParsedPersonRow]


def _cell_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    if not s or s in {"-", "—", "–", "None"}:
        return None
    # Excel often stores IDs as floats / ints
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".", 1)[0]
    return s


def _match_field(header: str, label_names: Optional[List[str]] = None) -> Optional[str]:
    if not header:
        return None
    # Exact match against company label display names first
    if label_names:
        for ln in label_names:
            if _norm_header(ln) == header:
                return f"label:{ln}"
    # Exact / contained alias match, longest first
    candidates: List[Tuple[int, str]] = []
    for field_name, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            a = _norm_header(alias)
            if not a:
                continue
            if header == a or a in header or header in a:
                candidates.append((len(a), field_name))
    if not candidates:
        # Fuzzy: header contained in a label name or vice versa
        if label_names:
            for ln in label_names:
                n = _norm_header(ln)
                if n and (header == n or n in header or header in n):
                    return f"label:{ln}"
        return None
    candidates.sort(reverse=True)
    return candidates[0][1]


def _score_mapping(mapping: Dict[str, int]) -> int:
    score = 0
    weights = {
        "personal_number": 5,
        "first_name": 4,
        "last_name": 4,
        "full_name": 5,
        "phone": 2,
        "pakal": 3,
        "role": 2,
        "department": 1,
        "squad": 1,
        "platoon": 1,
        "rank": 1,
    }
    for k, w in weights.items():
        if k in mapping:
            score += w
    for k in mapping:
        if k.startswith("label:"):
            score += 2
    # Prefer first+last over ambiguous "שם" alone
    if "first_name" in mapping and "last_name" in mapping:
        score += 3
    return score


def _find_header_row(
    ws, max_scan: int = 15, label_names: Optional[List[str]] = None
) -> Tuple[int, Dict[str, int], Dict[str, str]]:
    best = (0, {}, {})  # score, mapping col_index, display mapping
    max_row = min(max_scan, ws.max_row or 0)
    max_col = min(40, ws.max_column or 0)
    for r in range(1, max_row + 1):
        mapping: Dict[str, int] = {}
        display: Dict[str, str] = {}
        for c in range(1, max_col + 1):
            raw = ws.cell(r, c).value
            header = _norm_header(raw)
            field_name = _match_field(header, label_names)
            if field_name and field_name not in mapping:
                mapping[field_name] = c
                display[field_name] = str(raw).strip() if raw is not None else field_name
        score = _score_mapping(mapping)
        # Need at least a name signal
        has_name = (
            "full_name" in mapping
            or ("first_name" in mapping and "last_name" in mapping)
            or "first_name" in mapping
        )
        if has_name and score > best[0]:
            best = (score, mapping, display)
    return best[0], best[1], best[2]  # type: ignore[return-value]


def _infer_role_name(pakal: Optional[str], role_col: Optional[str], squad: Optional[str]) -> str:
    blob = " ".join(x for x in (role_col, pakal, squad) if x).lower()
    blob_compact = blob.replace('"', "").replace("״", "").replace("׳", "")
    if any(h in blob or h in blob_compact for h in COMMANDER_HINTS):
        return "מפקד"
    if any(h in blob or h in blob_compact for h in NCO_HINTS):
        return "מפקד זוטר"
    if squad and "מפקד" in squad:
        return "מפקד"
    return "חייל"


def _split_qualifications(pakal: Optional[str], role_name: str) -> List[str]:
    if not pakal:
        return []
    parts = re.split(r"[/,|\\]+", pakal)
    out: List[str] = []
    for p in parts:
        name = p.strip()
        if not name:
            continue
        # Role words are not qualifications
        n = _norm_header(name)
        if n in {"מפקד", "חייל", "מפקד זוטר"}:
            continue
        if name not in out:
            out.append(name)
    return out


def _build_notes(
    platoon: Optional[str],
    department: Optional[str],
    squad: Optional[str],
    extra: Optional[str] = None,
) -> Optional[str]:
    bits = []
    if platoon:
        bits.append(f"פלוגה: {platoon}")
    if department:
        bits.append(f"מחלקה: {department}")
    if squad:
        bits.append(f"כיתה: {squad}")
    if extra:
        bits.append(extra)
    return " · ".join(bits) if bits else None


def _split_multi_values(raw: Optional[str]) -> List[str]:
    if not raw:
        return []
    parts = re.split(r"[/,|\\]+", raw)
    out: List[str] = []
    for p in parts:
        name = p.strip()
        if name and name not in out:
            out.append(name)
    return out


def _parse_rows(
    ws,
    mapping: Dict[str, int],
    start_row: int,
    label_names: Optional[List[str]] = None,
) -> List[ParsedPersonRow]:
    rows: List[ParsedPersonRow] = []
    max_row = ws.max_row or 0
    for r in range(start_row + 1, max_row + 1):
        def get(field: str) -> Optional[str]:
            col = mapping.get(field)
            if not col:
                return None
            return _cell_str(ws.cell(r, col).value)

        first = get("first_name")
        last = get("last_name")
        full = get("full_name")
        if first or last:
            full_name = " ".join(x for x in (first, last) if x).strip()
        else:
            full_name = (full or "").strip()
        if not full_name:
            continue
        # Skip footer / totals
        if _norm_header(full_name) in {"סהכ", "סך הכל", "total"}:
            continue

        personal = get("personal_number")
        phone = get("phone")
        pakal = get("pakal")
        role_col = get("role")
        platoon = get("platoon")
        department = get("department")
        squad = get("squad")
        rank = get("rank")

        role_name = _infer_role_name(pakal, role_col, squad)
        quals = _split_qualifications(pakal, role_name)
        notes = _build_notes(platoon, department, squad, rank and f"דרגה: {rank}")

        label_values: Dict[str, List[str]] = {}
        names_by_norm = {_norm_header(n): n for n in (label_names or [])}

        for key in mapping:
            if not key.startswith("label:"):
                continue
            label_name = key[len("label:") :]
            col = mapping[key]
            cell = _cell_str(ws.cell(r, col).value)
            vals = _split_multi_values(cell)
            if vals:
                label_values[label_name] = vals

        alias_map = {
            "department": department,
            "squad": squad,
            "platoon": platoon,
        }
        alias_targets = {
            "department": ("מחלקה",),
            "squad": ("כיתה",),
            "platoon": ("פלוגה", "יחידה"),
        }
        for field, raw in alias_map.items():
            if not raw:
                continue
            matched: Optional[str] = None
            for target in alias_targets[field]:
                if _norm_header(target) in names_by_norm:
                    matched = names_by_norm[_norm_header(target)]
                    break
            if matched and matched not in label_values:
                label_values[matched] = _split_multi_values(raw)

        warnings: List[str] = []
        if not personal:
            warnings.append("ללא מספר אישי — יותאם לפי שם")
        rows.append(
            ParsedPersonRow(
                full_name=full_name,
                personal_number=personal,
                phone=phone,
                role_name=role_name,
                qualification_names=quals,
                label_values=label_values,
                notes=notes,
                warnings=warnings,
                raw={
                    "first_name": first,
                    "last_name": last,
                    "pakal": pakal,
                    "role": role_col,
                },
            )
        )
    return rows


def parse_people_workbook(
    data: bytes,
    preferred_sheet: Optional[str] = None,
    label_names: Optional[List[str]] = None,
) -> ParseResult:
    wb = load_workbook(io.BytesIO(data), data_only=True, keep_vba=False)
    sheet_options = [
        n
        for n in wb.sheetnames
        if not any(n.startswith(p) or n == p for p in SKIP_SHEET_PREFIXES)
        and not n.startswith("תרשים")
    ]
    if not sheet_options:
        sheet_options = list(wb.sheetnames)

    candidates: List[Tuple[int, str, Dict[str, int], Dict[str, str], int]] = []
    for name in sheet_options:
        ws = wb[name]
        score, mapping, display = _find_header_row(ws, label_names=label_names)
        if score <= 0:
            continue
        bonus = 0
        for hint in PREFERRED_SHEET_HINTS:
            if hint in name:
                bonus += 10 if name == hint else 4
        if preferred_sheet and name == preferred_sheet:
            bonus += 50
        # Prefer sheets that include personal number + pakal (מתוכנן)
        if "personal_number" in mapping:
            bonus += 5
        if "pakal" in mapping:
            bonus += 4
        candidates.append((score + bonus, name, mapping, display, score))

    if not candidates:
        raise ValueError(
            "לא נמצאה שורת כותרות מזוהה (שם / מספר אישי / טלפון). בדקו את הקובץ."
        )

    candidates.sort(key=lambda x: x[0], reverse=True)
    _, sheet_name, mapping, display, _ = candidates[0]
    ws = wb[sheet_name]

    # Re-find header row index for start
    header_row = 1
    max_row = min(15, ws.max_row or 0)
    max_col = min(40, ws.max_column or 0)
    for r in range(1, max_row + 1):
        trial: Dict[str, int] = {}
        for c in range(1, max_col + 1):
            raw = ws.cell(r, c).value
            field_name = _match_field(_norm_header(raw), label_names)
            if field_name and field_name not in trial:
                trial[field_name] = c
        if trial == mapping:
            header_row = r
            break

    rows = _parse_rows(ws, mapping, header_row, label_names=label_names)
    if not rows:
        raise ValueError(f"בגיליון «{sheet_name}» לא נמצאו שורות עם שמות חיילים")

    # Deduplicate by personal_number then full_name (keep first richer)
    deduped: List[ParsedPersonRow] = []
    seen_pn: set = set()
    seen_name: set = set()
    for row in rows:
        if row.personal_number:
            key = row.personal_number
            if key in seen_pn:
                continue
            seen_pn.add(key)
        else:
            key = row.full_name.strip().lower()
            if key in seen_name:
                continue
            seen_name.add(key)
        deduped.append(row)

    return ParseResult(
        sheet_name=sheet_name,
        sheet_options=sheet_options,
        column_mapping=display,
        rows=deduped,
    )


def _role_by_name(roles: List[Role], wanted: str) -> Role:
    wanted_n = _norm_header(wanted)
    for r in roles:
        if _norm_header(r.name) == wanted_n:
            return r
    # fuzzy contains
    for r in roles:
        if wanted_n in _norm_header(r.name) or _norm_header(r.name) in wanted_n:
            return r
    return roles[0]


def _get_or_create_qualification(
    db: Session, company_id: int, name: str, cache: Dict[str, Qualification]
) -> Tuple[Qualification, bool]:
    key = _norm_header(name)
    if key in cache:
        return cache[key], False
    existing = (
        db.query(Qualification)
        .filter(Qualification.company_id == company_id, Qualification.name == name)
        .first()
    )
    if existing:
        cache[key] = existing
        return existing, False
    # case-insensitive match
    for q in (
        db.query(Qualification).filter(Qualification.company_id == company_id).all()
    ):
        if _norm_header(q.name) == key:
            cache[key] = q
            return q, False
    q = Qualification(company_id=company_id, name=name.strip())
    db.add(q)
    db.flush()
    cache[key] = q
    return q, True


def apply_people_import(
    db: Session,
    company_id: int,
    parsed: ParseResult,
    *,
    commit: bool,
) -> Tuple[List[dict], int, int, int, int, List[str]]:
    """Returns preview dicts + counts. If commit=False, does not persist creates."""
    roles = (
        db.query(Role)
        .filter(Role.company_id == company_id, Role.is_active.is_(True))
        .all()
    )
    if not roles:
        raise ValueError("אין תפקידים מוגדרים בפלוגה — הוסיפו תפקיד בהגדרות לפני הייבוא")

    people = (
        db.query(Person)
        .options(joinedload(Person.qualifications))
        .filter(Person.company_id == company_id)
        .all()
    )
    by_pn = {
        p.personal_number: p
        for p in people
        if p.personal_number
    }
    by_name = {p.full_name.strip().lower(): p for p in people}

    qual_cache: Dict[str, Qualification] = {
        _norm_header(q.name): q
        for q in db.query(Qualification)
        .filter(Qualification.company_id == company_id)
        .all()
    }

    labels = load_company_labels(db, company_id)
    labels_by_norm = {_norm_header(lb.name): lb for lb in labels if lb.is_active}
    option_cache: Dict[tuple, object] = {}

    created = updated = skipped = quals_created = 0
    warnings: List[str] = []
    preview_rows: List[dict] = []

    def resolve_label_preview(row_labels: Dict[str, List[str]]) -> Dict[str, List[str]]:
        out: Dict[str, List[str]] = {}
        for name, vals in (row_labels or {}).items():
            lb = labels_by_norm.get(_norm_header(name))
            if not lb or not vals:
                continue
            mode = lb.selection_mode
            is_single = (
                mode == LabelSelectionMode.SINGLE or str(mode) == "single"
            )
            use_vals = vals[:1] if is_single else vals
            out[lb.name] = use_vals
        return out

    def apply_labels_to_person(person: Person, row_labels: Dict[str, List[str]]) -> None:
        resolved = resolve_label_preview(row_labels)
        if not resolved:
            return
        # Replace only labels present in the row (leave others untouched)
        touched_ids = set()
        for name, vals in resolved.items():
            lb = labels_by_norm[_norm_header(name)]
            touched_ids.add(lb.id)
            db.query(PersonLabelAssignment).filter(
                PersonLabelAssignment.person_id == person.id,
                PersonLabelAssignment.label_id == lb.id,
            ).delete(synchronize_session=False)
            for v in vals:
                opt = get_or_create_option(db, lb, v, option_cache)  # type: ignore[arg-type]
                db.add(
                    PersonLabelAssignment(
                        person_id=person.id,
                        label_id=lb.id,
                        option_id=opt.id,
                    )
                )

    for row in parsed.rows:
        match: Optional[Person] = None
        if row.personal_number and row.personal_number in by_pn:
            match = by_pn[row.personal_number]
        elif row.full_name.strip().lower() in by_name:
            match = by_name[row.full_name.strip().lower()]

        role = _role_by_name(roles, row.role_name or "חייל")
        action = "update" if match else "create"
        label_preview = resolve_label_preview(row.label_values)

        if not commit:
            preview_rows.append(
                {
                    "full_name": row.full_name,
                    "personal_number": row.personal_number,
                    "phone": row.phone,
                    "role_name": role.name,
                    "qualification_names": row.qualification_names,
                    "label_names": label_preview,
                    "notes": row.notes,
                    "action": action,
                    "match_person_id": match.id if match else None,
                    "warnings": row.warnings,
                }
            )
            if action == "create":
                created += 1
            else:
                updated += 1
            continue

        qual_ids: List[int] = []
        for qn in row.qualification_names:
            q, was_new = _get_or_create_qualification(db, company_id, qn, qual_cache)
            if was_new:
                quals_created += 1
            qual_ids.append(q.id)

        if match:
            match.full_name = row.full_name
            match.role_id = role.id
            if row.personal_number:
                match.personal_number = row.personal_number
            if row.phone:
                match.phone = row.phone
            if row.notes:
                match.notes = row.notes
            db.query(PersonQualification).filter(
                PersonQualification.person_id == match.id
            ).delete()
            for qid in qual_ids:
                db.add(PersonQualification(person_id=match.id, qualification_id=qid))
            apply_labels_to_person(match, row.label_values)
            updated += 1
            action = "update"
            pid = match.id
        else:
            person = Person(
                company_id=company_id,
                full_name=row.full_name,
                role_id=role.id,
                personal_number=row.personal_number,
                phone=row.phone,
                notes=row.notes,
            )
            db.add(person)
            db.flush()
            for qid in qual_ids:
                db.add(PersonQualification(person_id=person.id, qualification_id=qid))
            apply_labels_to_person(person, row.label_values)
            by_name[person.full_name.strip().lower()] = person
            if person.personal_number:
                by_pn[person.personal_number] = person
            created += 1
            action = "create"
            pid = person.id

        preview_rows.append(
            {
                "full_name": row.full_name,
                "personal_number": row.personal_number,
                "phone": row.phone,
                "role_name": role.name,
                "qualification_names": row.qualification_names,
                "label_names": label_preview,
                "notes": row.notes,
                "action": action,
                "match_person_id": pid,
                "warnings": row.warnings,
            }
        )

    if commit:
        db.commit()

    return preview_rows, created, updated, skipped, quals_created, warnings
